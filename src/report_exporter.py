# src/report_exporter.py
import os
import io
from datetime import datetime
from typing import Optional
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

class ReportExporter:
    """报告导出器，支持PDF和Word格式"""
    
    def __init__(self, output_dir: str = "reports"):
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)
    
    def export_to_pdf(self, report_content: str, session_id: str) -> str:
        """导出为PDF格式"""
        filename = f"{self.output_dir}/report_{session_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        
        doc = SimpleDocTemplate(filename, pagesize=letter)
        styles = getSampleStyleSheet()
        story = []
        
        # 添加标题
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=16,
            alignment=1,  # 居中
            spaceAfter=30
        )
        story.append(Paragraph("校园消费生态智能访谈报告", title_style))
        story.append(Spacer(1, 12))
        
        # 处理报告内容，按行分割
        lines = report_content.split('\n')
        for line in lines:
            if line.startswith('#'):
                level = line.count('#')
                text = line.lstrip('#').strip()
                if level == 1:
                    story.append(Paragraph(text, styles['Heading1']))
                elif level == 2:
                    story.append(Paragraph(text, styles['Heading2']))
                else:
                    story.append(Paragraph(text, styles['Heading3']))
            elif line.strip() and not line.startswith('|'):
                story.append(Paragraph(line, styles['Normal']))
            elif line.startswith('|'):
                # 简单处理表格
                story.append(Paragraph(line, styles['Code']))
            story.append(Spacer(1, 6))
        
        doc.build(story)
        return filename

    def export_to_pdf_bytes(self, report_content: str) -> bytes:
        """导出PDF字节流（用于浏览器下载）。"""
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=letter)
        styles = getSampleStyleSheet()
        story = []

        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=16,
            alignment=1,
            spaceAfter=30
        )
        story.append(Paragraph("校园消费生态智能访谈报告", title_style))
        story.append(Spacer(1, 12))

        lines = report_content.split('\n')
        for line in lines:
            if line.startswith('#'):
                level = line.count('#')
                text = line.lstrip('#').strip()
                if level == 1:
                    story.append(Paragraph(text, styles['Heading1']))
                elif level == 2:
                    story.append(Paragraph(text, styles['Heading2']))
                else:
                    story.append(Paragraph(text, styles['Heading3']))
            elif line.strip() and not line.startswith('|'):
                story.append(Paragraph(line, styles['Normal']))
            elif line.startswith('|'):
                story.append(Paragraph(line, styles['Code']))
            story.append(Spacer(1, 6))

        doc.build(story)
        return buf.getvalue()
    
    def export_to_word(self, report_content: str, session_id: str) -> str:
        """导出为Word格式"""
        filename = f"{self.output_dir}/report_{session_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
        
        doc = Document()
        
        # 添加标题
        title = doc.add_heading('校园消费生态智能访谈报告', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # 添加时间
        doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph()
        
        # 处理报告内容
        lines = report_content.split('\n')
        for line in lines:
            if line.startswith('##'):
                doc.add_heading(line.lstrip('#').strip(), level=2)
            elif line.startswith('#'):
                doc.add_heading(line.lstrip('#').strip(), level=1)
            elif line.strip():
                doc.add_paragraph(line)
        
        doc.save(filename)
        return filename

    def export_to_word_bytes(self, report_content: str) -> bytes:
        """导出Word字节流（用于浏览器下载）。"""
        doc = Document()

        title = doc.add_heading('校园消费生态智能访谈报告', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph()

        lines = report_content.split('\n')
        for line in lines:
            if line.startswith('##'):
                doc.add_heading(line.lstrip('#').strip(), level=2)
            elif line.startswith('#'):
                doc.add_heading(line.lstrip('#').strip(), level=1)
            elif line.strip():
                doc.add_paragraph(line)

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()